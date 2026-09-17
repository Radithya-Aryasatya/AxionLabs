"""tools/build_dock_twins.py
Builds REAL digital-twin layouts for Dock 2 and Dock 3 from your Excel files.

Inputs (your files):
  Possible Dock 2 Space Utilization.xlsx -> sheet 'Dock 2 - Demo' -> Dock 2
  Possible Dock 3 Sequence.xlsx          -> sheet 'Dock 2 - Demo' -> Dock 3

Outputs (overwrite the 3-fake-crate placeholders):
  assets/mock_docks/mock_layout_dock2.json
  assets/mock_docks/mock_layout_dock3.json

Packing: deterministic first-fit shelf packer (heavy floor-first, fragile
top-last - same loading-priority idea as Dock 1 in app.py). Produces real
non-overlapping x/y/z positions inside the same 2.4 x 2.4 x 6.0 m truck
Dock 1 uses, so the twin UI renders identically.

Run:  python AxionLabs/tools/build_dock_twins.py
"""
import json
import os
from itertools import permutations

import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Swap-folder pipeline: drop a newer manifest Excel into
# AxionLabs/dock_manifests/ (VSCODE only) and the twin rebuilds on next
# refresh. Falls back to the original Downloads files when the swap
# folder copy is absent.
_MANIFEST_DIR = os.path.join(BASE_DIR, "dock_manifests")
_DOWNLOADS = r"c:/Users/rian.dhanisaputra/Downloads"
DOCK2_XLSX = os.path.join(_DOWNLOADS, "Possible Dock 2 Space Utilization.xlsx")
DOCK3_XLSX = os.path.join(_DOWNLOADS, "Possible Dock 3 Sequence.xlsx")
_SWAP2 = os.path.join(_MANIFEST_DIR, "dock2_manifest.xlsx")
_SWAP3 = os.path.join(_MANIFEST_DIR, "dock3_manifest.xlsx")
if os.path.isfile(_SWAP2):
    DOCK2_XLSX = _SWAP2
if os.path.isfile(_SWAP3):
    DOCK3_XLSX = _SWAP3
MOCK_DIR = os.path.join(BASE_DIR, "assets", "mock_docks")

TRUCK_W_M, TRUCK_H_M, TRUCK_D_M = 2.4, 2.4, 6.0
TRUCK_W, TRUCK_H, TRUCK_D = 240.0, 240.0, 600.0


def read_demo_sheet(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Dock 2 - Demo"] if "Dock 2 - Demo" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    out = []
    for r in rows[1:]:
        if r[idx["Package ID"]] is None:
            continue
        out.append({
            "package_id": str(r[idx["Package ID"]]),
            "tracking": str(r[idx["Tracking Number"]]),
            "description": str(r[idx["Item Description"]]),
            "quantity": int(r[idx["Box Quantity"]]),
            "weight": float(r[idx["Box Weight (kg)"]]),
            "L": float(r[idx["Length (cm)"]]),
            "W": float(r[idx["Width (cm)"]]),
            "H": float(r[idx["Height (cm)"]]),
            "fragile": str(r[idx["Fragile"]]).strip().lower() == "yes",
            "handling": str(r[idx["Handling Instructions"]]),
            "zone": str(r[idx["Storage Zone"]]),
            "sequence": int(r[idx["Unloading Sequence"]]),
            "city": str(r[idx["Destination City"]]),
        })
    return out

def shelf_pack(rows):
    units = []
    for r in rows:
        name = f"{r['package_id']} {r['description']}"
        for _ in range(r["quantity"]):
            units.append({"sku": name, "row": r,
                          "dims": (r["L"], r["W"], r["H"]),
                          "weight": r["weight"], "fragile": r["fragile"]})
    heavy = [u for u in units if not u["fragile"] and u["weight"] > 15]
    light = [u for u in units if not u["fragile"] and u["weight"] <= 15]
    fragile = [u for u in units if u["fragile"]]
    for grp in (heavy, light, fragile):
        grp.sort(key=lambda u: -(u["dims"][0] * u["dims"][1] * u["dims"][2]))
    ordered = heavy + light + fragile
    shelves = []
    packed, unfitted = [], []
    for u in ordered:
        placed = False
        rots = sorted(set(permutations(u["dims"])),
                      key=lambda d: (d[1], d[2], d[0]))
        for dims in rots:
            dx, dy, dz = dims
            if dx > TRUCK_W or dy > TRUCK_H or dz > TRUCK_D:
                continue
            for sh in shelves:
                if dy > sh["h"]:
                    continue
                if u["fragile"] and sh["z0"] < 60 and len(shelves) > 1:
                    continue
                for row in sh["rows"]:
                    if dz <= row["depth"] and row["x"] + dx <= TRUCK_W:
                        packed.append((u, (row["x"], sh["z0"], row["y0"]), dims))
                        row["x"] += dx
                        placed = True
                        break
                if placed:
                    break
                used = sum(r2["depth"] for r2 in sh["rows"])
                if used + dz <= TRUCK_D and dx <= TRUCK_W:
                    sh["rows"].append({"y0": used, "x": dx, "depth": dz})
                    packed.append((u, (0, sh["z0"], used), dims))
                    placed = True
                    break
            if placed:
                break
            top = sum(s["h"] for s in shelves)
            if top + dy <= TRUCK_H:
                shelves.append({"z0": top, "h": dy,
                                "rows": [{"y0": 0, "x": dx, "depth": dz}]})
                packed.append((u, (0, top, 0), dims))
                placed = True
                break
        if not placed:
            unfitted.append(u)
    return packed, unfitted


def build_layout(dock_number, rows):
    packed, unfitted = shelf_pack(rows)
    packed_items, counters = [], {}
    for u, pos, dims in packed:
        n = counters.get(u["sku"], 0) + 1
        counters[u["sku"]] = n
        r = u["row"]
        packed_items.append({
            "name": f"{u['sku']} #{n}",
            "part_number": f"{r['package_id']}-{n}",
            "position": [float(pos[0]), float(pos[1]), float(pos[2])],
            "dimensions": [float(dims[0]), float(dims[1]), float(dims[2])],
            "weight": float(r["weight"]),
            "fragile": bool(r["fragile"]),
            "package_id": r["package_id"],
            "zone": r["zone"],
            "sequence": r["sequence"],
            "city": r["city"],
            "handling": r["handling"],
        })
    by_sku = {}
    for r in rows:
        by_sku[f"{r['package_id']} {r['description']}"] = r
    summary = []
    for sku, r in by_sku.items():
        n_packed = counters.get(sku, 0)
        summary.append({
            "name": sku, "package_id": r["package_id"],
            "description": r["description"], "quantity": r["quantity"],
            "packed": n_packed, "remaining": r["quantity"] - n_packed,
            "fragile": r["fragile"], "max_load": 10 if r["fragile"] else 100,
            "weight": r["weight"], "dimensions": [r["L"], r["W"], r["H"]],
            "zone": r["zone"], "sequence": r["sequence"], "city": r["city"],
            "handling": r["handling"], "tracking": r["tracking"],
        })
    truck_vol = TRUCK_W * TRUCK_H * TRUCK_D
    packed_vol = sum(p["dimensions"][0] * p["dimensions"][1]
                     * p["dimensions"][2] for p in packed_items)
    total_expected = sum(r["quantity"] for r in rows)
    return {
        "id": f"Fleet Monitoring | Dock {dock_number}",
        "truck_name": f"Truck-Dock{dock_number}",
        "truck_dimensions": [TRUCK_W_M, TRUCK_H_M, TRUCK_D_M],
        "fill_percentage": round(packed_vol / truck_vol * 100, 1),
        "loading_in_progress": True,
        "doors_closing": False,
        "truck_moving": False,
        "manifest": summary,
        "packed_items": packed_items,
        "manifest_summary": summary,
        "total_items_expected": total_expected,
        "packed_count": len(packed_items),
        "unfitted_count": len(unfitted),
        "unfitted_detail": [{"name": u["sku"],
                             "package_id": u["row"]["package_id"]}
                            for u in unfitted],
        "status": "LOADING",
        "gemini_analysis": {},
        "anomaly_history": [],
    }


def main():
    os.makedirs(MOCK_DIR, exist_ok=True)
    for dock_number, xlsx in [(2, DOCK2_XLSX), (3, DOCK3_XLSX)]:
        rows = read_demo_sheet(xlsx)
        layout = build_layout(dock_number, rows)
        path = os.path.join(MOCK_DIR, f"mock_layout_dock{dock_number}.json")
        with open(path, "w") as f:
            json.dump(layout, f, indent=2)
        print(f"Dock {dock_number} done: packed {layout['packed_count']}/"
              f"{layout['total_items_expected']} fill "
              f"{layout['fill_percentage']}% -> {path}")


if __name__ == "__main__":
    main()


