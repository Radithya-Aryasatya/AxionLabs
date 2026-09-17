"""
services/invoice_generator.py
=============================
Builds the downloadable "Load & Packing Invoice" Word document (.docx) for
a dock's fleet — the printable counterpart of the manager's
"📋 Mark as Finished" action.

- build_invoice_docx(fleet, doc_id=None) -> bytes
    Pure function (no streamlit dependency). EVERY number and table row is
    derived live from the fleet's packing plan / manifest — nothing is
    hardcoded. Package ID / Destination City come straight from the worker
    import (or dock manifests); missing optional values degrade to "—"
    instead of crashing or fabricating. Handling Instructions are removed
    by design everywhere — the manifest table shows a Fragile flag instead.
- invoice_doc_id / invoice_filename
    Helpers used by the manager controls (doc ID pattern AX-D{dock}-YYYYMMDD-NN).

Layout mirrors the approved example document: centered title block, meta
strip, AT A GLANCE KPI tiles, current load state, truck profile, volumetric
utilization bar, cargo manifest table, and the unloading sequence chain.
Colors were sampled directly from that document.
"""

import io
import math
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

try:
    from docx.oxml import OxmlElement
except ImportError:  # python-docx >= 1.2 moved the element factory
    from docx.oxml.parser import OxmlElement
from docx.oxml.ns import qn

# --- palette sampled from the approved example document ---------------------
NAVY = "12233F"        # primary text / headers / table header fill
SLATE = "64748B"       # labels, sub-lines
TEAL = "0F766E"        # accent: status, utilization fill, KPI highlights
AMBER = "B45309"       # fragile-cargo highlight
PANEL = "F4F7FA"       # light panel/tile fill
AMBER_PANEL = "FFF4E5"  # fragile tile fill
TRACK = "E8EAED"       # utilization bar track / light table borders
BODY = "1F2937"        # table body text
WHITE = "FFFFFF"

MIME_DOCX = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

_INVOICE_COMPANY = "PT. Kawan Lama Sejahtera"
_INVOICE_TITLE = "LOAD & PACKING INVOICE"
_DASH = "—"


# --- low-level docx styling helpers -----------------------------------------


def _shade(cell, fill):
    """Apply a background fill to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _table_borders(table, color=TRACK, sz=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def _tile_gaps(table):
    """White 2pt borders so adjacent shaded tiles read as separated cards."""
    _table_borders(table, color=WHITE, sz=16)


def _cell_margins(table, top=60, bottom=60, left=80, right=80):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("bottom", bottom),
                     ("left", left), ("right", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _fixed_layout(table):
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)


def _runs(par, runs, before=0, after=0):
    """Fill a paragraph with styled runs: (text, size_pt, bold, color_hex)."""
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    for text, size, bold, color in runs:
        run = par.add_run(str(text))
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "Calibri"
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return par


def _doc_par(doc, runs, before=0, after=0):
    return _runs(doc.add_paragraph(), runs, before=before, after=after)


def _cell_par(cell, runs, first=True, after=0):
    par = cell.paragraphs[0] if first and cell.paragraphs else cell.add_paragraph()
    return _runs(par, runs, after=after)


def _section_header(doc, title, subtitle=None, before=14):
    _doc_par(doc, [(title, 13, True, NAVY)], before=before, after=1)
    if subtitle:
        _doc_par(doc, [(subtitle, 7.5, False, SLATE)], after=6)


# --- live data collection ----------------------------------------------------


def _to_f(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _resolve_fragile(row, m):
    """Fragile resolution across the two manifest shapes in this app.

    Mock/Excel-driven manifests carry an explicit ``fragile`` flag. Dock 1's
    worker import does not: there the flag is encoded as ``max_load == box
    weight`` (non-fragile cargo gets ``max_load = inf``).
    """
    if m.get("fragile") is not None:
        return bool(m["fragile"])
    for src in (row, m):
        if src.get("fragile") is True:
            return True
    ml = _to_f(m.get("max_load", row.get("max_load")))
    w = _to_f(m.get("weight", row.get("weight")))
    if ml is None or not w or (isinstance(ml, float) and math.isinf(ml)):
        return False
    return ml <= w + 1e-6


def _description_of(merged):
    """Description without the leading package-id token (example format)."""
    desc = merged.get("description") or merged.get("item_description")
    if not desc:
        desc = str(merged.get("name", ""))
        pkg = merged.get("package_id")
        if pkg and str(desc).startswith(str(pkg)):
            desc = str(desc)[len(str(pkg)):].strip()
    return str(desc)


def _manifest_rows(plan, packed_items, manifest):
    """Normalize the dock's manifest into invoice rows (sorted by sequence)."""
    m_by_name = {str(m.get("name")): m for m in manifest}
    summary = list(plan.get("manifest_summary") or [])
    sources = summary if summary else manifest

    rows = []
    for src in sources:
        m = m_by_name.get(str(src.get("name")), {})
        merged = {**m, **src}
        name = str(merged.get("name", ""))
        if not name:
            continue
        qty = int(_to_f(merged.get("quantity")) or 0)
        packed = int(_to_f(merged.get("packed")) or 0)
        if "packed" not in merged:
            packed = sum(
                1 for p in packed_items
                if str(p.get("name", "")).split("#")[0].strip() == name
            )
        wpu = _to_f(merged.get("weight"))
        if wpu is None:
            weights = [
                _to_f(p.get("weight")) for p in packed_items
                if str(p.get("name", "")).split("#")[0].strip() == name
                and _to_f(p.get("weight")) is not None
            ]
            wpu = sum(weights) / len(weights) if weights else None
        rows.append({
            "sequence": merged.get("sequence"),
            "package_id": merged.get("package_id"),
            "description": _description_of(merged),
            "quantity": max(0, qty),
            "packed": max(0, packed),
            "weight": wpu,
            "destination": merged.get("city") or merged.get("destination"),
            "fragile": _resolve_fragile(merged, m),
        })

    def _seq_key(r):
        s = _to_f(r["sequence"])
        return (1e9, 0.0) if s is None else (s, 0.0)

    rows.sort(key=_seq_key)
    return rows


def _collect(fleet):
    """Pull every invoice number live from the fleet's packing data."""
    plan = fleet.packing_layout or {}
    layout = plan.get("layout", {}) or {}
    packed_items = layout.get("packed_items", []) or []
    rows = _manifest_rows(plan, packed_items, list(fleet.manifest or []))

    total_expected = int(_to_f(plan.get("total_items_expected")) or 0) or sum(
        r["quantity"] for r in rows)
    packed_count = int(_to_f(plan.get("packed_count")) or 0) or (
        len(packed_items) if packed_items else sum(r["packed"] for r in rows))

    planned_weight = sum(
        r["quantity"] * r["weight"] for r in rows if r["weight"] is not None)
    packed_weight = sum(_to_f(p.get("weight")) or 0 for p in packed_items)
    if not packed_items:
        packed_weight = sum(
            r["packed"] * r["weight"] for r in rows if r["weight"] is not None)

    tw, th, td = (fleet.truck_dimensions or (0, 0, 0))[:3]
    truck_vol = (_to_f(tw) or 0) * (_to_f(th) or 0) * (_to_f(td) or 0)
    packed_vol = 0.0
    for p in packed_items:
        dims = p.get("dimensions") or []
        if len(dims) == 3 and all(_to_f(x) is not None for x in dims):
            packed_vol += (_to_f(dims[0]) * _to_f(dims[1])
                           * _to_f(dims[2])) / 1e6  # cm³ -> m³

    fill_field = _to_f(plan.get("fill_percentage", fleet.fill_percentage)) or 0.0
    if truck_vol > 0 and packed_vol > 0:
        fill = packed_vol / truck_vol * 100.0
    else:
        fill = fill_field
        if truck_vol > 0 and fill > 0 and packed_vol <= 0:
            packed_vol = truck_vol * fill / 100.0

    fragile_count = sum(r["quantity"] for r in rows if r["fragile"])
    status_value = getattr(fleet.status, "value", str(fleet.status))
    status_display = ("FINISHED" if status_value == "INSPECTED - CLEAR"
                      else status_value)

    return {
        "rows": rows, "plan": plan,
        "total_expected": total_expected, "packed_count": packed_count,
        "planned_weight": planned_weight, "packed_weight": packed_weight,
        "truck_vol": truck_vol, "packed_vol": packed_vol,
        "fill": fill, "fragile_count": fragile_count,
        "status_display": status_display,
        "seq_priority": plan.get("sequence_priority") or "ON",
    }


def _utilization_png(pct):
    """Teal-on-track utilization bar as PNG (matches the example's gauge).

    Returns None when matplotlib is unavailable — the caller falls back to a
    text bar so the invoice is never blocked by an optional dependency.
    """
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return None
    pct = max(0.0, min(100.0, pct))
    fig, ax = plt.subplots(figsize=(6.2, 0.42), dpi=200)
    fig.patch.set_alpha(0.0)
    ax.barh([0], [100], color="#" + TRACK, height=0.62)
    ax.barh([0], [pct], color="#" + TEAL, height=0.62)
    if pct >= 12:
        ax.text(pct / 2, 0, f"{pct:.1f}%", ha="center", va="center",
                color="white", fontsize=9, fontweight="bold")
    ax.set_xlim(0, 100)
    ax.axis("off")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight", transparent=True)
    plt.close(fig)
    buf.seek(0)
    return buf


def _text_bar(pct):
    filled = int(round(max(0.0, min(100.0, pct)) / 5))
    return "█" * filled + "░" * (20 - filled)


def build_invoice_docx(fleet, doc_id=None):
    """Build the Load & Packing Invoice .docx for a fleet. Returns bytes."""
    d = _collect(fleet)
    rows = d["rows"]
    dock = fleet.dock_number
    truck = fleet.truck_name or f"Truck-Dock{dock}"
    if doc_id is None:
        doc_id = invoice_doc_id(dock)
    finished = d["status_display"] == "FINISHED"
    accent = TEAL if finished else AMBER

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.55)
    sec.left_margin = sec.right_margin = Inches(0.62)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(8)
    normal.font.color.rgb = RGBColor.from_string(BODY)

    # --- title block ---
    _doc_par(doc, [(_INVOICE_COMPANY, 24, True, NAVY)], before=6, after=0)
    _doc_par(doc, [(_INVOICE_TITLE, 11.5, True, SLATE)], after=8)

    # --- meta strip (DOCUMENT ID / DOCK / TRUCK / FLEET STATUS) ---
    meta = doc.add_table(rows=2, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    _tile_gaps(meta)
    _cell_margins(meta)
    _fixed_layout(meta)
    meta_labels = ["DOCUMENT ID", "DOCK", "TRUCK", "FLEET STATUS"]
    meta_values = [doc_id, f"Dock {dock}", truck, d["status_display"]]
    for i, col in enumerate(meta.columns):
        for cell in col.cells:
            cell.width = Inches(1.815)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _shade(cell, PANEL)
        _cell_par(col.cells[0], [(meta_labels[i], 7, True, SLATE)], after=1)
        _cell_par(col.cells[1], [
            (str(meta_values[i]), 11.5, True, NAVY if i < 3 else accent),
        ])

    # --- AT A GLANCE KPI tiles ---
    _section_header(doc, "AT A GLANCE", "Executive Load Snapshot")
    kpis = [
        ("PACKAGES", f"{d['packed_count']:,}",
         (f"{d['packed_count']:,} packages packed / "
          f"{d['total_expected']:,} from manifest"), TEAL, PANEL),
        ("LOADED WEIGHT", f"{d['packed_weight']:,.0f} kg",
         f"{d['planned_weight']:,.0f} kg planned total", NAVY, PANEL),
        ("VOLUME UTILIZATION", f"{d['fill']:.1f}%",
         (f"{d['packed_vol']:.2f} m³ packed / "
          f"{d['truck_vol']:.2f} m³ truck"), TEAL, PANEL),
        ("FRAGILE CARGO", f"{d['fragile_count']:,}",
         "Amount of objects requiring care", AMBER, AMBER_PANEL),
    ]
    kpi = doc.add_table(rows=1, cols=4)
    kpi.alignment = WD_TABLE_ALIGNMENT.CENTER
    _tile_gaps(kpi)
    _cell_margins(kpi)
    _fixed_layout(kpi)
    for i, (label, value, sub, vcolor, fill) in enumerate(kpis):
        cell = kpi.rows[0].cells[i]
        cell.width = Inches(1.815)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade(cell, fill)
        _cell_par(cell, [(label, 7, True, SLATE)], after=1)
        _cell_par(cell, [(value, 19, True, vcolor)], after=1, first=False)
        _cell_par(cell, [(sub, 6.5, False, SLATE)], first=False, after=0)

    # --- current load state ---
    _section_header(doc, "CURRENT LOAD STATE")
    _doc_par(doc, [(d["status_display"], 19, True, accent)], after=1)
    _doc_par(doc, [(f"Sequence priority: {d['seq_priority']}", 8, False, SLATE)],
             after=2)

    # --- truck profile ---
    _section_header(doc, "TRUCK PROFILE")
    tw, th, td = (fleet.truck_dimensions or (0, 0, 0))[:3]
    for label, value in (
        ("Truck ID", truck),
        ("Internal dimensions", f"{_to_f(tw) or 0:.1f} × {_to_f(th) or 0:.1f} "
                                f"× {_to_f(td) or 0:.1f} m"),
        ("Packed cargo volume", f"{d['packed_vol']:.2f} m³"),
        ("Packed cargo weight", f"{d['packed_weight']:,.0f} kg"),
    ):
        _doc_par(doc, [
            (f"{label}:  ", 8, False, SLATE),
            (str(value), 8, True, NAVY),
        ], after=1)

    # --- volumetric utilization ---
    _section_header(doc, "VOLUMETRIC UTILIZATION")
    png = _utilization_png(d["fill"])
    if png is not None:
        doc.add_picture(png, width=Inches(6.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        _doc_par(doc, [(_text_bar(d["fill"]), 10, True, TEAL)], after=2)
    headroom = max(0.0, 100.0 - d["fill"])
    _doc_par(doc, [
        (f"{d['fill']:.1f}%", 8, True, TEAL),
        (" occupied  •  ", 8, False, SLATE),
        (f"{headroom:.1f}%", 8, True, SLATE),
        (" nominal volumetric headroom", 8, False, SLATE),
    ], after=2)

    # --- cargo manifest table ---
    _section_header(
        doc, "CARGO MANIFEST",
        "Destination sequence, package quantities, weight, and fragile flags",
    )
    headers = ["Sequence", "Package ID", "Description", "Manifest", "Packed",
               "Weight/Unit", "Destination", "Fragile"]
    widths = [0.55, 0.8, 2.46, 0.62, 0.6, 0.78, 0.85, 0.6]
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_borders(tbl, TRACK, 4)
    _cell_margins(tbl, top=40, bottom=40, left=60, right=60)
    _fixed_layout(tbl)
    for i, cell in enumerate(tbl.rows[0].cells):
        cell.width = Inches(widths[i])
        _shade(cell, NAVY)
        _cell_par(cell, [(headers[i], 7.5, True, WHITE)], after=0)
    for r in rows:
        seq_s = _to_f(r["sequence"])
        vals = [
            f"{int(seq_s)}" if seq_s is not None else _DASH,
            str(r["package_id"] or _DASH),
            r["description"],
            f"{r['quantity']:,}",
            f"{r['packed']:,}",
            (f"{r['weight']:g} kg" if r["weight"] is not None else _DASH),
            str(r["destination"] or _DASH),
            "Yes" if r["fragile"] else _DASH,
        ]
        cells = tbl.add_row().cells
        for i, text in enumerate(vals):
            cells[i].width = Inches(widths[i])
            color = (AMBER if (text == "Yes" and headers[i] == "Fragile")
                     else BODY)
            _cell_par(cells[i], [(text, 7.5, False, color)], after=0)

    # --- unloading sequence chain ---
    _section_header(doc, "UNLOADING SEQUENCE")
    chain_runs = []
    for i, r in enumerate(rows):
        s = _to_f(r["sequence"])
        num = f"{int(s):02d}" if s is not None else "--"
        label = (r["destination"] if r["destination"] not in (None, "", _DASH)
                 else r["description"])
        chain_runs.append((f"{num}  {label}", 8, True, NAVY))
        if i < len(rows) - 1:
            chain_runs.append(("   →   ", 8, True, TEAL))
    if chain_runs:
        _doc_par(doc, chain_runs, after=2)
    else:
        _doc_par(doc, [("No packed cargo to sequence.", 8, False, SLATE)])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def invoice_doc_id(dock_number, seq=1, now=None):
    """Document ID in the example's pattern: AX-D{dock}-YYYYMMDD-{seq}."""
    stamp = (now or datetime.now()).strftime("%Y%m%d")
    return f"AX-D{dock_number}-{stamp}-{seq:02d}"


def invoice_filename(fleet, now=None):
    stamp = (now or datetime.now()).strftime("%Y%m%d")
    return f"AxionLabs_Invoice_Dock{fleet.dock_number}_{stamp}.docx"