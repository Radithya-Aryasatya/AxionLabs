"""
verify_rejection_flow.py
========================
End-to-end validation of the "unfitted package reason" chain:

1. reclassify_floating_items() in app.py — extracts the real function body
   via AST (app.py has streamlit side effects at import, so we exec just the
   function) and verifies unstable items are moved packed -> unfitted with
   the footprint_instability reason.
2. Engine reason tags — packs an overweight box with py3dbp and verifies
   the engine writes rejection_reason='overweight'.
3. build_fleet_state — verifies the fleet's packing_layout carries
   reason/detail and per-type 'unfitted' counts in manifest_summary.
4. Invoice — builds the .docx for a fleet WITH rejections and asserts the
   "REJECTED / NOT LOADED PACKAGES" section, reason labels, and the packed X
   of Y summary line; also asserts the all-packed case prints the honest
   "All manifest items packed and loaded." line.
5. Mock docks — loads dock2/dock3 layout JSONs through the mock factory and
   verifies reasons survive into the fleet.

Run:  python verify_rejection_flow.py   (from the AxionLabs folder)
"""

import ast
import io
import json
import math
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from py3dbp import Packer, Bin, Item, REJECTION_UNSTABLE
from state.fleet_state import build_fleet_from_packing_result
from state.fleet_state import Fleet, FleetStatus
from services.invoice_generator import build_invoice_docx

FAILURES = []


def check(name, cond, detail=""):
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}")
    if not cond:
        FAILURES.append(f"{name} {('— ' + str(detail)) if detail else ''}")


def _doc_text(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


# --------------------------------------------------------------------------
print("1) reclassify_floating_items (app.py helper, extracted via AST)")
# --------------------------------------------------------------------------
_app_src = open(os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "app.py"),
    encoding="utf-8").read()
_tree = ast.parse(_app_src)
_fn = next(n for n in ast.walk(_tree)
           if isinstance(n, ast.FunctionDef)
           and n.name == "reclassify_floating_items")
_ns = {"REJECTION_UNSTABLE": "footprint_instability"}
exec(compile(ast.Module(body=[_fn], type_ignores=[]), "<app.py>", "exec"),
     _ns)
reclassify_floating_items = _ns["reclassify_floating_items"]


class FakeItem:
    def __init__(self, name, partno):
        self.name, self.partno = name, partno


class FakeBin:
    def __init__(self, items):
        self.items = list(items)
        self.unfitted_items = []
        self.unfitted_reasons = {}


it_a, it_b, it_c = (FakeItem(f"Box #{i}", f"Box #{i}") for i in (1, 2, 3))
bin_obj = FakeBin([it_a, it_b, it_c])
moved = reclassify_floating_items(bin_obj, ["Box #2"])
check("returns moved count", moved == 1, moved)
check("moved item out of bin.items",
      [i.name for i in bin_obj.items] == ["Box #1", "Box #3"])
check("moved item into unfitted_items",
      [i.name for i in bin_obj.unfitted_items] == ["Box #2"])
check("reason tagged footprint_instability",
      it_b.rejection_reason == REJECTION_UNSTABLE)
check("detail filled",
      "75%" in it_b.rejection_detail, it_b.rejection_detail)
check("bin.unfitted_reasons updated",
      bin_obj.unfitted_reasons.get("Box #2") == REJECTION_UNSTABLE)
check("idempotent (no double-move)",
      reclassify_floating_items(bin_obj, ["Box #2"]) == 0
      and len(bin_obj.unfitted_items) == 1)
check("empty floating list is a no-op",
      reclassify_floating_items(bin_obj, []) == 0)


# --------------------------------------------------------------------------
print("2) engine writes 'overweight' reason on rejected item")
# --------------------------------------------------------------------------
packer = Packer()
engine_bin = Bin("Truck-Test", (100, 100, 100), 10)  # capacity 10 kg
packer.addBin(engine_bin)
packer.addItem(Item("Light #1", "Light", "box", (10, 10, 10), 1.0, 1, 100, True, "white"))
packer.addItem(Item("Heavy #1", "Heavy", "box", (10, 10, 10), 99.0, 1, 100, True, "white"))
packer.pack(distribute_items=False)
unf = engine_bin.unfitted_items
check("heavy item rejected",
      len(unf) == 1 and unf[0].partno == "Heavy #1",
      [(i.name, i.partno) for i in unf])
check("reason == overweight",
      getattr(unf[0], "rejection_reason", None) == "overweight"
      or engine_bin.unfitted_reasons.get("Heavy #1") == "overweight",
      getattr(unf[0], "rejection_reason", None))


# --------------------------------------------------------------------------
print("3) build_fleet_from_packing_result carries reasons + unfitted counts")
# --------------------------------------------------------------------------
manifest = [
    {"name": "Light", "quantity": 1, "max_load": 100, "fragile": False,
     "package_id": "PKG-L", "city": "Rotterdam"},
    {"name": "Heavy", "quantity": 1, "max_load": 100, "fragile": False,
     "package_id": "PKG-H", "city": "Atlanta"},
]
fleet = build_fleet_from_packing_result(
    manifest=manifest, packer=packer,
    truck_w=1.0, truck_h=1.0, truck_d=1.0, truck_name="Truck-T",
)
layout = fleet.packing_layout["layout"]
check("unfitted item present in fleet layout",
      len(layout["unfitted_items"]) == 1)
check("fleet layout reason == overweight",
      layout["unfitted_items"][0]["reason"] == "overweight",
      layout["unfitted_items"][0]["reason"])
check("unfitted_count == 1", fleet.packing_layout["unfitted_count"] == 1)
summary = {s["name"]: s for s in fleet.packing_layout["manifest_summary"]}
check("per-row unfitted count (Heavy=1)",
      summary["Heavy"].get("unfitted") == 1, summary["Heavy"])
check("per-row unfitted count (Light=0)",
      summary["Light"].get("unfitted") == 0)
check("packed_count == 1", fleet.packing_layout["packed_count"] == 1)


# --------------------------------------------------------------------------
print("4) invoice shows REJECTED / NOT LOADED PACKAGES section")
# --------------------------------------------------------------------------
docx = build_invoice_docx(fleet, doc_id="AX-D1-20260918-99")
text = _doc_text(docx)
check("section header present", "REJECTED / NOT LOADED PACKAGES" in text)
check("human reason label in doc",
      "Exceeds truck weight capacity" in text)
check("package id in rejected table", "PKG-H" in text)
check("summary line 'Packed 1 of 2'",
      "Packed 1 of 2 packages — 1 rejected by constraints." in text)

# all-packed fleet -> honest "nothing rejected" line
plan_ok = {
    "layout": {"packed_items": [
        {"name": "Light #1", "weight": 1.0, "dimensions": [10, 10, 10]},
    ], "unfitted_items": []},
    "manifest_summary": [
        {"name": "Light", "package_id": "PKG-L", "quantity": 1,
         "total_expected": 1, "packed": 1, "remaining": 0,
         "unfitted": 0, "fragile": False, "max_load": 100,
         "city": "Rotterdam", "weight": 1.0},
    ],
    "total_items_expected": 1, "packed_count": 1, "unfitted_count": 0,
    "fill_percentage": 1.0, "sequence_priority": "ON",
}
fleet_ok = Fleet(id="TEST-OK", dock_number=1,
                 truck_dimensions=(1.0, 1.0, 1.0), manifest=manifest,
                 packing_layout=plan_ok, status=FleetStatus.INSPECTED_CLEAR,
                 fill_percentage=1.0, truck_name="Truck-T", source="mock")
text_ok = _doc_text(build_invoice_docx(fleet_ok, doc_id="AX-D1-20260918-98"))
check("all-packed case prints honest line",
      "All manifest items packed and loaded." in text_ok)
check("no rejected table when nothing rejected",
      "rejected by constraints." not in text_ok)


# --------------------------------------------------------------------------
print("5) mock dock layouts carry reasons into the fleet")
# --------------------------------------------------------------------------
from services.mock_fleet_factory import _build_fleet_from_layout

for dock, expected_reason in ((2, "overweight"),
                              (3, "footprint_instability")):
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "assets", "mock_docks", f"mock_layout_dock{dock}.json")
    data = json.load(open(path, encoding="utf-8"))
    f, _ = _build_fleet_from_layout(data)
    unf = f.packing_layout["layout"]["unfitted_items"]
    reasons = {u.get("reason") for u in unf}
    check(f"dock{dock}: all unfitted have reason == {expected_reason}",
          unf and reasons == {expected_reason}, reasons)
    check(f"dock{dock}: unfitted_count matches list",
          f.packing_layout["unfitted_count"] == len(unf))

print()
if FAILURES:
    print(f"{len(FAILURES)} FAILURE(S):")
    for f_ in FAILURES:
        print(" -", f_)
    sys.exit(1)
print("ALL CHECKS PASSED")
